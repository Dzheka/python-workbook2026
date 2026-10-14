"""Build the "Many Under One Name" PowerPoint deck for the lists topic.

Starts from the row of numbered variables a list replaces, then walks the
material in ``lists/_docs``: creating lists, indexing and slicing, operators
and methods, iteration, `split`/`join`, copies versus references, and nested
lists -- ending with where to start in the list exercises.

The drawing primitives are the same ones used by the earlier decks, so the
whole course looks like one set.

Run:  uv run --with python-pptx python slides/lists/build_lists_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("lists.pptx")

# ---------------------------------------------------------------- palette --
NAVY = RGBColor(0x1E, 0x2A, 0x38)
BLUE = RGBColor(0x37, 0x76, 0xAB)
YELLOW = RGBColor(0xFF, 0xD4, 0x3B)
GRAY = RGBColor(0x5A, 0x66, 0x72)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_FG = RGBColor(0xE8, 0xEE, 0xF4)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)

SANS = "Trebuchet MS"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)  # left/right margin
BODY_TOP = Inches(1.75)
BODY_W = W - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

_section = {"n": 0, "name": ""}


# ------------------------------------------------------------- primitives --
def _box(
    slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None
):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def _tf(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(
    tf,
    text,
    size,
    color,
    bold=False,
    font=SANS,
    first=False,
    space_before=0,
    space_after=0,
    align=PP_ALIGN.LEFT,
    italic=False,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if text:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = color
    return p


def _rich(p, parts, size, default_color, font=SANS):
    """parts: list of (text, color|None, bold, mono)."""
    for text, color, bold, mono in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = MONO if mono else font
        f.color.rgb = color or default_color


def _markup(text):
    """`code` -> mono+blue, *bold* -> bold navy."""
    parts, buf, i = [], "", 0
    while i < len(text):
        ch = text[i]
        if ch in "`*":
            end = text.find(ch, i + 1)
            if end > i:
                if buf:
                    parts.append((buf, None, False, False))
                    buf = ""
                inner = text[i + 1 : end]
                if ch == "`":
                    parts.append((inner, BLUE, False, True))
                else:
                    parts.append((inner, NAVY, True, False))
                i = end + 1
                continue
        buf += ch
        i += 1
    if buf:
        parts.append((buf, None, False, False))
    return parts


def slide(title=None, kicker=None, footer=True):
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=WHITE)
    if title is not None:
        _box(s, M, Inches(0.62), Inches(0.09), Inches(0.72), fill=YELLOW)
        tf = _tf(s, M + Inches(0.26), Inches(0.5), BODY_W - Inches(0.26), Inches(1.0))
        if kicker:
            _para(tf, kicker.upper(), 12, BLUE, bold=True, first=True, space_after=2)
            _para(tf, title, 30, NAVY, bold=True)
        else:
            _para(tf, title, 30, NAVY, bold=True, first=True)
    if footer and _section["name"]:
        tf = _tf(s, M, H - Inches(0.52), BODY_W, Inches(0.3))
        _para(tf, _section["name"], 10, GRAY, first=True)
    return s


def bullets(s, items, x=None, y=None, w=None, size=18, gap=9):
    """items: str, or (str, sub_level)."""
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    tf = _tf(s, x, y, w, H - y - Inches(0.7))
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        if text == "":
            _para(tf, "", 6, GRAY, first=first)
            first = False
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(0 if p is tf.paragraphs[0] else gap)
        p.space_after = Pt(0)
        p.level = lvl
        p.alignment = PP_ALIGN.LEFT
        glyph = "▸  " if lvl == 0 else "        –  "
        r = p.add_run()
        r.text = glyph
        r.font.size = Pt(size - (2 if lvl else 0))
        r.font.name = SANS
        r.font.color.rgb = BLUE if lvl == 0 else GRAY
        r.font.bold = lvl == 0
        _rich(p, _markup(text), size - (2 if lvl else 0), NAVY if lvl == 0 else GRAY)
    return tf


def code(s, lines, x=None, y=None, w=None, size=15, caption=None, dark=True):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.46).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=NAVY if dark else LIGHT,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.2), w - Inches(0.4), h - Inches(0.3))
    for i, ln in enumerate(lines):
        col = CODE_FG if dark else NAVY
        if ln.strip().startswith("#"):
            col = RGBColor(0x8E, 0xA6, 0xBD) if dark else GRAY
        _para(
            tf, ln if ln else " ", size, col, font=MONO, first=(i == 0), space_after=0
        )
    if caption:
        ctf = _tf(s, x, y + h + Inches(0.06), w, Inches(0.3))
        _para(ctf, caption, 13, GRAY, first=True, italic=True)
    return y + h + (Inches(0.42) if caption else Inches(0.24))


def out(s, lines, x=None, y=None, w=None, size=15, label="Output"):
    x = M if x is None else x
    w = BODY_W if w is None else w
    lh = Pt(size * 1.42).emu
    h = Emu(int(lh * len(lines)) + Inches(0.62).emu)
    _box(
        s,
        x,
        y,
        w,
        h,
        fill=LIGHT,
        line=RGBColor(0xD5, 0xDE, 0xE6),
        shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.045,
    )
    tf = _tf(s, x + Inches(0.22), y + Inches(0.16), w - Inches(0.4), h - Inches(0.26))
    _para(tf, label.upper(), 10, GRAY, bold=True, first=True, space_after=3)
    for ln in lines:
        _para(tf, ln if ln else " ", size, NAVY, font=MONO)
    return y + h + Inches(0.24)


def table(
    s,
    headers,
    rows,
    x=None,
    y=None,
    w=None,
    size=14,
    col_w=None,
    mono_cols=(),
    height=None,
):
    x = M if x is None else x
    y = BODY_TOP if y is None else y
    w = BODY_W if w is None else w
    h = height or Inches(0.42 + 0.36 * len(rows))
    shp = s.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    tbl = shp.table
    tbl.first_row = True
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.margin_left = c.margin_right = Inches(0.1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(c.text_frame, htxt, size, WHITE, bold=True, first=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            mono = j in mono_cols
            _rich(
                p, _markup(val) if not mono else [(val, BLUE, False, True)], size, NAVY
            )
    return y + h + Inches(0.2)


def note(s, text, y, kind="tip"):
    colors = {
        "tip": (BLUE, "TIP"),
        "warn": (RED, "WATCH OUT"),
        "ok": (GREEN, "REMEMBER"),
    }
    col, label = colors[kind]
    h = Inches(0.78)
    _box(s, M, y, BODY_W, h, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    _box(s, M, y, Inches(0.07), h, fill=col)
    tf = _tf(
        s,
        M + Inches(0.26),
        y + Inches(0.1),
        BODY_W - Inches(0.5),
        h - Inches(0.2),
        anchor=MSO_ANCHOR.MIDDLE,
    )
    _para(tf, label, 10, col, bold=True, first=True, space_after=2)
    p = tf.add_paragraph()
    _rich(p, _markup(text), 15, NAVY)
    return y + h + Inches(0.2)


def section(name, subtitle, points):
    _section["n"] += 1
    _section["name"] = name
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.22), W, Inches(0.22), fill=YELLOW)
    tf = _tf(s, M, Inches(2.1), Inches(8.4), Inches(3.0))
    _para(tf, f"PART {_section['n']}", 16, YELLOW, bold=True, first=True, space_after=8)
    _para(tf, name, 44, WHITE, bold=True, space_after=10)
    _para(tf, subtitle, 19, RGBColor(0xA9, 0xBD, 0xD1))
    tf2 = _tf(s, Inches(9.0), Inches(2.3), Inches(3.6), Inches(3.0))
    _para(tf2, "IN THIS PART", 11, YELLOW, bold=True, first=True, space_after=10)
    for pt in points:
        _para(tf2, "•  " + pt, 15, RGBColor(0xD6, 0xE2, 0xEC), space_after=7)
    return s


# ================================================================== TITLE ==
def title_slide():
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.28), W, Inches(0.28), fill=YELLOW)
    _box(s, M, Inches(1.5), Inches(0.12), Inches(1.5), fill=YELLOW)
    tf = _tf(s, M + Inches(0.4), Inches(1.4), Inches(9.5), Inches(2.0))
    _para(tf, "PYTHON · TOPIC 5", 15, YELLOW, bold=True, first=True, space_after=12)
    _para(tf, "Many Under One Name", 52, WHITE, bold=True, space_after=10)
    _para(
        tf,
        "Lists: indexing, slicing, methods, iteration and copies",
        22,
        RGBColor(0xA9, 0xBD, 0xD1),
    )
    tf2 = _tf(s, M + Inches(0.4), Inches(4.3), Inches(10.5), Inches(2.0))
    _para(
        tf2,
        "A variable holds one value.",
        18,
        RGBColor(0x8F, 0xB6, 0xD9),
        first=True,
        space_after=8,
    )
    _para(
        tf2,
        "A list holds as many as you need — in order.",
        18,
        RGBColor(0x8F, 0xB6, 0xD9),
    )
    tf3 = _tf(s, M + Inches(0.4), Inches(6.2), Inches(10.0), Inches(0.5))
    _para(
        tf3,
        "python-workbook2026  ·  lists",
        13,
        RGBColor(0x6E, 0x8A, 0xA3),
        first=True,
    )


def roadmap_slide():
    s = slide("What we cover today", kicker="Roadmap", footer=False)
    table(
        s,
        ["Part", "Topic", "You will be able to"],
        [
            ["1", "*Why lists*", "create a list and see why it beats `n1, n2, n3`"],
            ["2", "*Indexing & slicing*", "read one item, or a whole piece"],
            ["3", "*Operations & methods*", "grow, shrink, sort and search a list"],
            ["4", "*Iteration*", "loop with `for`, `enumerate` and `zip`"],
            ["5", "*Strings ↔ lists*", "`split` a line and `join` it back"],
            ["6", "*Copies & references*", "know when two names share one list"],
            ["7", "*Nested lists*", "build and walk a matrix"],
        ],
        y=BODY_TOP,
        col_w=(0.8, 3.6, 6.8),
        size=16,
    )
    note(
        s,
        "Exercises are still *functions* the tests call — only now the "
        "arguments and the results are lists.",
        Inches(5.6),
        kind="ok",
    )


# =================================================================== PART 1 ==
def part1():
    section(
        "Why lists",
        "One name for many values",
        ["Five variables", "List literals", "`list()` and `range`"],
    )

    s = slide("Five scores, five variables", kicker="Without a list")
    y = code(
        s,
        [
            "score1 = 85",
            "score2 = 92",
            "score3 = 78",
            "score4 = 90",
            "score5 = 88",
            "avg = (score1 + score2 + score3",
            "       + score4 + score5) / 5",
        ],
        w=Inches(5.6),
    )
    code(
        s,
        [
            "scores = [85, 92, 78, 90, 88]",
            "",
            "avg = sum(scores) / len(scores)",
            "",
            "# a 6th score? one append, nothing else",
            "scores.append(95)",
        ],
        x=M + Inches(6.1),
        y=BODY_TOP,
        w=Inches(5.6),
    )
    note(
        s,
        "A list is *ordered* and *changeable*. The same code works for 5 "
        "scores or 5 000.",
        y + Inches(0.1),
        kind="ok",
    )

    s = slide("Creating lists", kicker="Literals")
    y = code(
        s,
        [
            "empty = []",
            "numbers = [2, 3, 5, 7, 11]",
            'fruits = ["apple", "banana", "cherry"]',
            'mixed = [1, "two", 3.0, True]       # allowed, rarely useful',
            "",
            "digits = list(range(10))            # [0, 1, ..., 9]",
            "letters = list(\"abc\")               # ['a', 'b', 'c']",
            "squares = [x * x for x in range(5)] # [0, 1, 4, 9, 16]",
        ],
    )
    note(
        s,
        "Square brackets, items separated by commas. `len(numbers)` tells "
        "you how many items there are.",
        y,
        kind="tip",
    )


# =================================================================== PART 2 ==
def part2():
    section(
        "Indexing & slicing",
        "Getting items out",
        ["`[i]` from 0", "Negative indices", "`[start:stop:step]`"],
    )

    s = slide("Indexes start at 0", kicker="Indexing")
    table(
        s,
        ["Item", '`"a"`', '`"b"`', '`"c"`', '`"d"`', '`"e"`'],
        [
            ["index", "`0`", "`1`", "`2`", "`3`", "`4`"],
            ["negative", "`-5`", "`-4`", "`-3`", "`-2`", "`-1`"],
        ],
        col_w=(2.2, 1.8, 1.8, 1.8, 1.8, 1.8),
        size=17,
    )
    y = code(
        s,
        [
            'letters = ["a", "b", "c", "d", "e"]',
            "letters[0]     # 'a'  - first",
            "letters[-1]    # 'e'  - last",
            "letters[len(letters) // 2]   # 'c' - middle",
            "letters[5]     # IndexError: list index out of range",
        ],
        y=Inches(3.15),
    )
    note(
        s,
        "The last index is `len(items) - 1`, not `len(items)`. `items[-1]` "
        "says the same thing, shorter.",
        y,
        kind="warn",
    )

    s = slide("Slicing takes a piece", kicker="Slicing")
    y = code(
        s,
        [
            "nums = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9]",
            "",
            "nums[2:5]      # [2, 3, 4]      - stop is NOT included",
            "nums[:3]       # [0, 1, 2]      - from the start",
            "nums[7:]       # [7, 8, 9]      - to the end",
            "nums[::2]      # [0, 2, 4, 6, 8] - every second",
            "nums[::-1]     # [9, 8, ..., 0] - reversed",
            "nums[8:20]     # [8, 9]         - no IndexError",
        ],
    )
    note(
        s,
        "A slice always returns a *new list*. Slicing past the end is fine — "
        "that is exactly what `get_page` needs.",
        y,
        kind="ok",
    )


# =================================================================== PART 3 ==
def part3():
    section(
        "Operations & methods",
        "Changing and asking a list",
        ["`+`, `*`, `in`", "Add and remove", "Sorting", "Method vs function"],
    )

    s = slide("Operators work on lists", kicker="Operations")
    table(
        s,
        ["Write", "Result", "Meaning"],
        [
            ["`[1, 2] + [3]`", "`[1, 2, 3]`", "concatenation — a new list"],
            ["`[0] * 3`", "`[0, 0, 0]`", "repetition"],
            ["`3 in [1, 2, 3]`", "`True`", "membership"],
            ["`len([1, 2, 3])`", "`3`", "number of items"],
            ["`sum`, `min`, `max`", "a number", "work on any list of numbers"],
            ["`[1, 2] < [1, 3]`", "`True`", "compared item by item"],
        ],
        col_w=(3.4, 2.6, 5.2),
        size=16,
    )

    s = slide("Methods change the list in place", kicker="Methods")
    table(
        s,
        ["Method", "What it does", "Returns"],
        [
            ["`append(x)`", "add `x` to the end", "`None`"],
            ["`insert(i, x)`", "put `x` at index `i`", "`None`"],
            ["`extend(other)`", "add every item of `other`", "`None`"],
            ["`remove(x)`", "delete the first `x` (ValueError if absent)", "`None`"],
            ["`pop()` / `pop(i)`", "take out the last / the `i`-th item", "that item"],
            ["`index(x)`, `count(x)`", "where is `x`, how many `x`", "a number"],
            ["`sort()`, `reverse()`", "reorder", "`None`"],
        ],
        col_w=(3.2, 5.6, 2.4),
        size=15,
    )
    note(
        s,
        "`append` adds *one* item: `a.append([4, 5])` puts a list inside a "
        "list. To add several, use `extend`.",
        Inches(5.6),
        kind="warn",
    )

    s = slide("sort() or sorted()?", kicker="The classic trap")
    y = code(
        s,
        [
            "nums = [3, 1, 2]",
            "result = nums.sort()",
            "print(result)",
            "print(nums)",
        ],
        w=Inches(5.6),
    )
    code(
        s,
        [
            "nums = [3, 1, 2]",
            "result = sorted(nums)",
            "print(result)",
            "print(nums)",
        ],
        x=M + Inches(6.1),
        y=BODY_TOP,
        w=Inches(5.6),
    )
    out(s, ["None", "[1, 2, 3]"], y=y, w=Inches(5.6))
    y = out(s, ["[1, 2, 3]", "[3, 1, 2]"], x=M + Inches(6.1), y=y, w=Inches(5.6))
    note(
        s,
        "A *method* changes the list and returns `None`; a *function* like "
        "`sorted` leaves it alone and returns a new one. `reverse=True` works "
        "for both.",
        y,
        kind="ok",
    )


# =================================================================== PART 4 ==
def part4():
    section(
        "Iteration",
        "Visiting every item",
        ["`for item in items`", "`enumerate`", "`zip`", "Building a new list"],
    )

    s = slide("Loop over items, not indexes", kicker="for")
    y = code(
        s,
        [
            'fruits = ["apple", "banana", "cherry"]',
            "",
            "for fruit in fruits:                  # the items themselves",
            "    print(fruit)",
            "",
            "for i, fruit in enumerate(fruits):    # index and item",
            "    print(i, fruit)",
            "",
            "for a, b in zip([1, 2, 3], [10, 20, 30]):   # two lists in step",
            "    print(a + b)",
        ],
    )
    note(
        s,
        "`for i in range(len(items))` works, but `enumerate` says what you "
        "mean. Reach for indexes only when you need them.",
        y,
        kind="tip",
    )

    s = slide("Accumulate into a new list", kicker="Pattern")
    y = code(
        s,
        [
            "def evens(numbers: list[int]) -> list[int]:",
            "    result = []",
            "    for n in numbers:",
            "        if n % 2 == 0:",
            "            result.append(n)",
            "    return result",
            "",
            "# the same as a list comprehension",
            "def evens(numbers: list[int]) -> list[int]:",
            "    return [n for n in numbers if n % 2 == 0]",
        ],
    )
    note(
        s,
        "Never `remove` from a list while you loop over it — items get "
        "skipped. Build a new list instead.",
        y,
        kind="warn",
    )


# =================================================================== PART 5 ==
def part5():
    section(
        "Strings ↔ lists",
        "Reading a line, writing a line",
        ["`split()`", "`join()`", "Converting the items"],
    )

    s = slide("split and join", kicker="Text in, text out")
    y = code(
        s,
        [
            'line = "3 14 15 92"',
            "parts = line.split()                 # ['3', '14', '15', '92']",
            "numbers = [int(p) for p in parts]    # [3, 14, 15, 92]",
            "",
            'words = ["red", "green", "blue"]',
            "\", \".join(words)                     # 'red, green, blue'",
            "\" \".join(str(n) for n in numbers)    # '3 14 15 92'",
        ],
    )
    note(
        s,
        "`join` is a method of the *separator* and needs strings: "
        '`", ".join([1, 2])` is a TypeError — convert with `str` first.',
        y,
        kind="warn",
    )


# =================================================================== PART 6 ==
def part6():
    section(
        "Copies & references",
        "Two names, one list?",
        ["`b = a` is not a copy", "Three ways to copy", "Lists in functions"],
    )

    s = slide("b = a does not copy", kicker="References")
    y = code(
        s,
        [
            "a = [1, 2, 3]",
            "b = a            # second name, same list",
            "b.append(4)",
            "print(a)",
            "",
            "c = a.copy()     # or a[:]  or list(a)",
            "c.append(5)",
            "print(a, c)",
        ],
        w=Inches(6.8),
    )
    out(
        s,
        ["[1, 2, 3, 4]", "[1, 2, 3, 4] [1, 2, 3, 4, 5]"],
        x=M + Inches(7.2),
        y=BODY_TOP,
        w=Inches(4.4),
    )
    note(
        s,
        "`a is b` asks *same object?*; `a == b` asks *same items?*. "
        "`copy_list` checks both.",
        y + Inches(0.1),
        kind="ok",
    )

    s = slide("A function can change your list", kicker="Side effects")
    y = code(
        s,
        [
            "def add_zero(items: list[int]) -> None:",
            "    items.append(0)          # changes the caller's list",
            "",
            "nums = [1, 2]",
            "add_zero(nums)",
            "print(nums)                  # [1, 2, 0]",
        ],
    )
    note(
        s,
        "When an exercise says *return a new list*, do not touch the one you "
        "were given — copy it, or build a fresh one.",
        y,
        kind="warn",
    )


# =================================================================== PART 7 ==
def part7():
    section(
        "Nested lists",
        "Lists of lists",
        ["Matrices", "`m[row][col]`", "The `* n` trap"],
    )

    s = slide("A matrix is a list of rows", kicker="2D")
    y = code(
        s,
        [
            "matrix = [",
            "    [1, 2, 3],",
            "    [4, 5, 6],",
            "    [7, 8, 9],",
            "]",
            "matrix[1][2]          # 6  - row 1, column 2",
            "",
            "for row in matrix:",
            '    print(" ".join(str(x) for x in row))',
        ],
        w=Inches(7.0),
    )
    out(s, ["1 2 3", "4 5 6", "7 8 9"], x=M + Inches(7.4), y=BODY_TOP, w=Inches(4.2))
    note(
        s,
        "First index picks the *row*, second the *column*. `len(matrix)` is "
        "the number of rows; `len(matrix[0])`, of columns.",
        y,
        kind="ok",
    )

    s = slide("Building one correctly", kicker="The * n trap")
    y = code(
        s,
        [
            "wrong = [[0] * 3] * 3            # three names for ONE row",
            "wrong[0][0] = 1",
            "print(wrong)   # [[1, 0, 0], [1, 0, 0], [1, 0, 0]]",
            "",
            "right = [[0] * 3 for _ in range(3)]   # three separate rows",
            "right[0][0] = 1",
            "print(right)   # [[1, 0, 0], [0, 0, 0], [0, 0, 0]]",
        ],
    )
    note(
        s,
        "`* n` on the *outer* list repeats a reference — the same trap as "
        "`b = a`. Build each row in a loop or comprehension.",
        y,
        kind="warn",
    )


# ========================================================= HOW IT'S GRADED ==
def grading():
    _section["name"] = "How the exercises work"
    s = slide("Where to start", kicker="Workflow")
    table(
        s,
        ["Level", "Exercises"],
        [
            ["warm-up", "`list_items`, `get_page`, `str_to_ints`, `reverse_list`"],
            ["loops", "`list_total`, `average`, `remove_outliers`, `divisors`"],
            ["copies & 2D", "`copy_list`, `to3x3matrix`, `to_matrix`"],
            ["algorithms", "`binary_search`, `merge_sort`, `sieve_of_eratosthenes`"],
            ["calculator", "`tokenize_expression` → … → `evaluate_postfix`"],
        ],
        col_w=(2.4, 8.8),
        size=16,
    )
    bullets(
        s,
        [
            "Same as functions: create `average.py` with the function from the statement",
            "`list_total` forbids `sum()` — the test reads your source",
            "`infix_to_postfix` imports your `precedence` — solve that one first",
        ],
        y=Inches(4.35),
        size=17,
    )
    out(
        s,
        ["uv run pytest lists/average/"],
        y=Inches(5.95),
        label="Check it",
    )


# ================================================================= CLOSING ==
def closing():
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=NAVY)
    _box(s, 0, H - Inches(0.28), W, Inches(0.28), fill=YELLOW)
    tf = _tf(s, Inches(1.3), Inches(1.5), Inches(10.7), Inches(3.6))
    _para(tf, "Your turn", 44, WHITE, bold=True, first=True, space_after=18)
    _para(
        tf,
        "Start with list_items and get_page, then work down the table of contents.",
        22,
        RGBColor(0xC3, 0xD3, 0xE1),
        space_after=26,
    )
    _para(
        tf,
        "Before you write a loop, ask: do I change this list, or return a new one?",
        17,
        RGBColor(0x8F, 0xB6, 0xD9),
    )
    tf2 = _tf(s, Inches(1.3), Inches(5.5), Inches(10.7), Inches(1.2))
    _para(
        tf2,
        "uv run pytest lists/",
        18,
        YELLOW,
        bold=True,
        first=True,
        space_after=10,
    )
    _para(tf2, "Good luck  ·  see you soon", 20, YELLOW, bold=True)


# ==================================================================== main ==
title_slide()
roadmap_slide()
part1()
part2()
part3()
part4()
part5()
part6()
part7()
grading()
closing()

prs.save(OUT)
print(f"{len(prs.slides)} slides -> {OUT}")
